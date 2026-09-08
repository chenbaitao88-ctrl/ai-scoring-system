"""Real file bytes, durable score facts and all public export aliases."""
import csv
import io
import json
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import test_result_derivation as rd
import test_review_case_store as rc
import test_score_attempt_store as sa
from routers import export, result_export, scores_admin
from services import authoritative_export_service as aes
from services.result_derivation_service import ResultDerivationService
from services.review_case_store import ReviewCaseStore
from services.score_attempt_store import ScoreAttemptStore

TASK = "11111111-1111-4111-8111-111111111111"
ITEM = "22222222-2222-4222-8222-222222222222"
ITEM2 = "33333333-3333-4333-8333-333333333333"
FORMATS = [("xlsx", "/api/export/excel"), ("docx", "/api/export/word"), ("csv", "/api/scores/export")]


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(rc, "TASK", TASK)
    monkeypatch.setattr(rc, "ITEM", ITEM)
    attempts = ScoreAttemptStore(tmp_path / "attempts")
    attempts.create_attempt(TASK, ITEM, sa.make_attempt(status="succeeded", task_id=TASK, item_id=ITEM))
    attempts.write_snapshot(TASK, ITEM, sa.make_snapshot())
    attempts.write_validation(TASK, ITEM, sa.make_validation())
    reviews = ReviewCaseStore(tmp_path / "reviews", attempt_store=attempts)
    derivation = ResultDerivationService(reviews, attempts, tmp_path)
    item = SimpleNamespace(item_id=ITEM, task_id=TASK, task_type="scoring_pipeline")
    task = SimpleNamespace(task_id=TASK, task_type="scoring_pipeline", batch_id="synthetic-batch", status="completed",
                           total_items=1, item_index=[item], revision=1)
    items = [item]
    manager = SimpleNamespace(get_task=lambda task_id: task if task_id == TASK else None,
                              list_items=lambda task_id: items,
                              list_tasks=lambda offset, limit: [task][offset:offset + limit])
    service = aes.AuthoritativeExportService(manager, derivation)
    app = FastAPI()
    app.include_router(result_export.router)
    app.include_router(export.router, prefix="/api/export")
    app.include_router(scores_admin.router, prefix="/api/scores")
    app.dependency_overrides[result_export.get_authoritative_export_service] = lambda: service
    app.dependency_overrides[result_export.get_result_derivation_service] = lambda: derivation
    with TestClient(app) as client:
        yield SimpleNamespace(attempts=attempts, reviews=reviews, derivation=derivation, manager=manager,
                              service=service, task=task, items=items, client=client, root=tmp_path)


def score_row(fmt, payload):
    if fmt == "xlsx":
        book = load_workbook(io.BytesIO(payload), data_only=False)
        assert book.sheetnames == ["导出说明", "权威结果", "维度分"]
        assert book["权威结果"]["C2"].data_type == "n"
        assert book["权威结果"]["H2"].data_type == "s"
        return dict(zip(aes.HEADERS, next(book["权威结果"].iter_rows(min_row=2, max_row=2, values_only=True))))
    if fmt == "docx":
        doc = Document(io.BytesIO(payload))
        assert "本次导出1条" in "".join(p.text for p in doc.paragraphs)
        return {row.cells[0].text: row.cells[1].text for row in doc.tables[0].rows}
    return next(csv.DictReader(io.StringIO(payload.decode("utf-8-sig"))))


@pytest.mark.parametrize("fmt,alias", FORMATS)
def test_formats_and_aliases_match_derived_score_and_ids(env, fmt, alias):
    rd.seed_adoption(env.reviews, env.attempts)
    expected = env.client.post(f"/api/result-export/tasks/{TASK}/items/{ITEM}/derive").json()
    for url, params in [(f"/api/result-export/tasks/{TASK}/export", {"format": fmt}), (alias, {"task_id": TASK})]:
        response = env.client.get(url, params=params)
        assert response.status_code == 200, response.text
        assert response.headers["cache-control"] == "no-store"
        row = score_row(fmt, response.content)
        assert float(row["总分"]) == expected["result"]["total_score"] == 82
        assert row["结果版本"] == expected["result"]["snapshot_id"]
        assert row["评分尝试"] == expected["result"]["attempt_id"]
        assert row["推导编号"] == expected["derivation_id"]
        assert row["采用记录"] == rc.ADP


@pytest.mark.parametrize("fmt,alias", FORMATS)
def test_adjusted_score_is_exported(env, fmt, alias):
    env.attempts.write_snapshot(TASK, ITEM, sa.make_snapshot(snapshot_id="snap-adjusted", total_score=85))
    env.attempts.write_validation(TASK, ITEM, sa.make_validation(validation_id="val-adjusted", snapshot_id="snap-adjusted"))
    rd.seed_adoption(env.reviews, env.attempts, snapshot_id="snap-adjusted", validation_id="val-adjusted")
    row = score_row(fmt, env.client.get(alias, params={"task_id": TASK}).content)
    assert float(row["总分"]) == 85
    assert row["结果版本"] == "snap-adjusted"


def test_lock_survives_new_successful_attempt_and_reopen_service(env):
    rd.seed_adoption(env.reviews, env.attempts)
    rd.seed_case(env.reviews)
    rd.seed_lock(env.reviews)
    env.attempts.create_attempt(TASK, ITEM, sa.make_attempt(
        attempt_id="attempt-new", task_id=TASK, item_id=ITEM, status="succeeded", attempt_number=2,
        previous_attempt_id=sa.ATTEMPT_ID, result_snapshot_ref="snapshot-new", validation_ref="validation-new"))
    env.attempts.write_snapshot(TASK, ITEM, sa.make_snapshot(snapshot_id="snapshot-new", attempt_id="attempt-new", total_score=91))
    env.attempts.write_validation(TASK, ITEM, sa.make_validation(validation_id="validation-new", attempt_id="attempt-new", snapshot_id="snapshot-new"))
    reopened_attempts = ScoreAttemptStore(env.root / "attempts")
    reopened_reviews = ReviewCaseStore(env.root / "reviews", attempt_store=reopened_attempts)
    service = aes.AuthoritativeExportService(env.manager, ResultDerivationService(reopened_reviews, reopened_attempts, env.root))
    for fmt, _ in FORMATS:
        row = score_row(fmt, service.export(TASK, fmt))
        assert float(row["总分"]) == 82
        assert row["确认方式"] == "manual_final_lock"
        assert row["锁定记录"] == "flk_demo_001"


@pytest.mark.parametrize("state", ["open", "assigned", "in_review", "waiting_for_evidence", "missing_snapshot", "corrupted_snapshot", "no_adoption"])
def test_blocked_state_never_returns_any_file(env, state):
    if state != "no_adoption":
        rd.seed_adoption(env.reviews, env.attempts)
    if state in ("open", "assigned", "in_review", "waiting_for_evidence"):
        rd.seed_case(env.reviews, status=state)
    elif state.endswith("snapshot"):
        path = env.attempts.root / "tasks" / TASK / "items" / ITEM / "snapshots" / f"{sa.SNAPSHOT_ID}.json"
        if state == "missing_snapshot":
            path.rename(path.with_suffix(".missing"))
        else:
            path.write_text("{broken")
    for fmt, alias in FORMATS:
        for url, params in [(alias, {"task_id": TASK}), (f"/api/result-export/tasks/{TASK}/export", {"format": fmt})]:
            response = env.client.get(url, params=params)
            assert response.status_code == 409
            assert response.json()["error"]["code"] == "EXPORT_BLOCKED_BY_REVIEW"
            assert "content-disposition" not in response.headers
    preview = env.client.get(f"/api/result-export/tasks/{TASK}/results").json()
    assert preview["items"][0]["exportable"] is False
    assert "result" not in preview["items"][0]


def test_full_task_does_not_silently_omit_blocked_item(env):
    rd.seed_adoption(env.reviews, env.attempts)
    second = SimpleNamespace(item_id=ITEM2, task_id=TASK, task_type="scoring_pipeline")
    env.items.append(second)
    env.task.item_index.append(second)
    env.task.total_items = 2
    url = f"/api/result-export/tasks/{TASK}/export"
    assert env.client.get(url).status_code == 409
    response = env.client.get(url, params={"format": "csv", "item_id": ITEM})
    assert response.status_code == 200
    row = score_row("csv", response.content)
    assert row["导出条数"] == "1" and row["任务总条数"] == "2"


@pytest.mark.parametrize("selection", [[ITEM, ITEM], [ITEM2], ["../unsafe"]])
def test_invalid_selection_rejected(env, selection):
    response = env.client.get(f"/api/result-export/tasks/{TASK}/export", params=[("item_id", value) for value in selection])
    assert response.status_code == 400


def test_missing_manifest_item_blocks_instead_of_silent_omission(env):
    env.items.clear()
    response = env.client.get(f"/api/result-export/tasks/{TASK}/export")
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "EXPORT_TASK_INCOMPLETE"


def test_empty_task_and_unknown_task(env):
    env.items.clear()
    env.task.item_index.clear()
    env.task.total_items = 0
    assert env.client.get(f"/api/result-export/tasks/{TASK}/export").status_code == 400
    assert env.client.get(f"/api/result-export/tasks/{ITEM2}/export").status_code == 404


def test_change_during_render_discards_file(env, monkeypatch):
    rd.seed_adoption(env.reviews, env.attempts)
    original = aes._csv
    def changed(selection):
        data = original(selection)
        env.task.revision += 1
        return data
    monkeypatch.setattr(aes, "_csv", changed)
    response = env.client.get(f"/api/result-export/tasks/{TASK}/export", params={"format": "csv"})
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "EXPORT_RESULT_CHANGED"


def test_preview_version_must_match_download(env):
    rd.seed_adoption(env.reviews, env.attempts)
    url = f"/api/result-export/tasks/{TASK}/export"
    body = {"format": "csv", "item_ids": [ITEM], "expected_derivations": {ITEM: "stale-version"}}
    response = env.client.post(url, json=body)
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "EXPORT_RESULT_CHANGED"
    body["expected_derivations"][ITEM] = env.service.preview(TASK)["items"][0]["derivation_id"]
    response = env.client.post(url, json=body)
    assert response.status_code == 200
    assert float(score_row("csv", response.content)["总分"]) == 82
    assert env.client.post(url, json={"format": "csv", "item_ids": []}).status_code == 422


def test_preview_and_exports_exclude_sensitive_fields(env):
    rd.seed_adoption(env.reviews, env.attempts)
    preview = env.client.get(f"/api/result-export/tasks/{TASK}/results")
    texts = [preview.text]
    for fmt, _ in FORMATS:
        payload = env.service.export(TASK, fmt)
        if fmt == "xlsx":
            book = load_workbook(io.BytesIO(payload))
            texts.append(str([[list(row) for row in sheet.values] for sheet in book]))
        elif fmt == "docx":
            doc = Document(io.BytesIO(payload))
            texts.append(str([[cell.text for row in table.rows for cell in row.cells] for table in doc.tables]))
        else:
            texts.append(payload.decode("utf-8-sig"))
    for text in texts:
        for secret in ("demo rationale", "input_fingerprint", "evidence_refs", "provider_demo", str(env.root)):
            assert secret not in text


def test_pagination_does_not_drop_later_scoring_tasks(env):
    evidence = SimpleNamespace(task_type="evidence_preparation_pipeline")
    tasks = [evidence] * 205 + [env.task]
    env.manager.list_tasks = lambda offset, limit: tasks[offset:offset + limit]
    assert env.service.list_tasks()[0]["task_id"] == TASK
