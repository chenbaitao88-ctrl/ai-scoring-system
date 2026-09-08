"""
SQLite数据库模型定义
使用SQLAlchemy ORM
"""
from datetime import datetime
from sqlalchemy import (
    Column, Integer, String, Float, DateTime, Text, Boolean,
    ForeignKey, JSON, create_engine, Index
)
from sqlalchemy.orm import declarative_base, relationship, sessionmaker
from config import DATABASE_URL, DATA_DIR

engine = create_engine(DATABASE_URL, connect_args={"check_same_thread": False})
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base = declarative_base()


def get_db():
    """获取数据库会话"""
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def init_db():
    """初始化数据库，创建所有表"""
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    # 检查数据库文件是否存在，但表不存在的情况
    db_path = DATA_DIR / "teams.db"
    if db_path.exists():
        import sqlite3
        conn = sqlite3.connect(str(db_path))
        cursor = conn.cursor()
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table'")
        tables = cursor.fetchall()
        conn.close()
        if not tables:  # 数据库文件存在但没有表，删除文件重新创建
            print("数据库文件存在但表不存在，删除并重新创建...")
            db_path.unlink()

    Base.metadata.create_all(bind=engine)
    # 迁移：添加新字段（如果不存在）
    _migrate_add_columns()

    # 自动初始化任务书
    _init_task_books()

    # GENERIC: 如果没有 competitions 记录，自动创建默认赛事
    _init_default_competition()


def _migrate_add_columns():
    """安全添加新列（SQLite兼容）"""
    import sqlite3
    db_path = DATA_DIR / "teams.db"
    if not db_path.exists():
        return

    conn = sqlite3.connect(str(db_path))
    cursor = conn.cursor()

    # 需要添加的新列
    new_columns = [
        # machine_scores 表的评论列
        ("machine_scores", "theme_comment", "TEXT"),
        ("machine_scores", "presentation_comment", "TEXT"),
        ("machine_scores", "process_comment", "TEXT"),
        ("machine_scores", "ai_literacy_comment", "TEXT"),
        ("machine_scores", "overall_comment", "TEXT"),
        # machine_scores 表的新功能列（2026-04-28）
        ("machine_scores", "model_name", "VARCHAR(100)"),
        ("machine_scores", "is_adopted", "BOOLEAN DEFAULT 0"),
        ("machine_scores", "scoring_session", "VARCHAR(100)"),
        ("machine_scores", "defense_questions", "JSON"),
        # P0-3 (2026-05-22): 锚点置信度与等级字段
        ("machine_scores", "confidence", "VARCHAR(20) DEFAULT 'UNKNOWN'"),
        ("works", "anchor_level", "VARCHAR(10) DEFAULT 'Lv0'"),
        # P0-5 (2026-05-22): Flag 规则引擎结果 (JSON, SQLite 实存 TEXT)
        ("machine_scores", "flags", "TEXT DEFAULT '[]'"),
        # P1-6 (2026-05-26): 校准后总分
        ("machine_scores", "calibrated_score", "FLOAT DEFAULT 0"),
        # teams 表的审核相关列
        ("teams", "status", "VARCHAR(20) DEFAULT 'confirmed'"),
        ("teams", "source", "VARCHAR(20) DEFAULT 'excel'"),
        # teams 表拆分指导老师字段（2026-05-12）
        ("teams", "teacher_name", "VARCHAR(50)"),
        ("teams", "teacher_phone", "VARCHAR(20)"),
        # === Phase 1.1: 双轨评分字段（2026-05-27）===
        # competitions 表
        ("competitions", "scoring_mode", "VARCHAR(20) DEFAULT 'mixed'"),
        ("competitions", "ai_weight", "FLOAT DEFAULT 0.4"),
        ("competitions", "human_weight", "FLOAT DEFAULT 0.6"),
        # machine_scores 表：AI 维度分
        ("machine_scores", "ai_theme_score", "FLOAT DEFAULT 0"),
        ("machine_scores", "ai_code_quality_score", "FLOAT DEFAULT 0"),
        ("machine_scores", "ai_completeness_score", "FLOAT DEFAULT 0"),
        ("machine_scores", "ai_aigc_score", "FLOAT DEFAULT 0"),
        ("machine_scores", "ai_total_score", "FLOAT DEFAULT 0"),
        # human_scores 表：人工维度分
        ("human_scores", "human_presentation_score", "FLOAT DEFAULT 0"),
        ("human_scores", "human_creativity_score", "FLOAT DEFAULT 0"),
        ("human_scores", "human_process_score", "FLOAT DEFAULT 0"),
        ("human_scores", "human_performance_score", "FLOAT DEFAULT 0"),
        ("human_scores", "human_total_score", "FLOAT DEFAULT 0"),
        # final_scores 表：综合分
        ("final_scores", "composite_score", "FLOAT DEFAULT 0"),
        ("final_scores", "scoring_mode", "VARCHAR(20) DEFAULT 'mixed'"),
        # Phase 2: AI 推断人工维度（ai_only 模式）
        ("machine_scores", "inferred_presentation_score", "FLOAT DEFAULT 0"),
        ("machine_scores", "inferred_creativity_score", "FLOAT DEFAULT 0"),
        ("machine_scores", "inferred_process_score", "FLOAT DEFAULT 0"),
        ("machine_scores", "inferred_performance_score", "FLOAT DEFAULT 0"),
        ("machine_scores", "inferred_total_score", "FLOAT DEFAULT 0"),
        # GENERIC: Phase 2 扩展 competitions 表字段
        ("competitions", "dimensions_config", "TEXT DEFAULT '{}'"),
        ("competitions", "naming_pattern", "VARCHAR(500)"),
        ("competitions", "naming_example", "VARCHAR(200)"),
        ("competitions", "judge_groups", "INTEGER DEFAULT 4"),
        ("competitions", "judges_per_group", "INTEGER DEFAULT 2"),
        ("competitions", "machine_weight", "FLOAT DEFAULT 0.6"),
        ("competitions", "task_book_ids", "TEXT DEFAULT '[]'"),
    ]

    for table, column, col_type in new_columns:
        try:
            cursor.execute(f"ALTER TABLE {table} ADD COLUMN {column} {col_type}")
        except Exception:
            pass  # 列已存在，忽略

    # 迁移teacher字段（格式：姓名（手机号））
    try:
        cursor.execute("SELECT id, teacher FROM teams WHERE teacher IS NOT NULL AND teacher != ''")
        import re
        for row in cursor.fetchall():
            team_id, teacher_value = row
            if teacher_value:
                # 格式：姓名（手机号）或 姓名 (手机号)，支持空格
                match = re.match(r'^(.+?)\s*[\(（](\d{11})[\)）]$', teacher_value)
                if match:
                    name = match.group(1).strip()
                    phone = match.group(2)
                else:
                    name = teacher_value
                    phone = ""

                cursor.execute(
                    "UPDATE teams SET teacher_name=?, teacher_phone=? WHERE id=?",
                    (name, phone, team_id)
                )
    except Exception as e:
        print(f"迁移teacher字段失败: {e}")
        pass

    conn.commit()
    conn.close()


def _init_task_books():
    """自动初始化任务书（如果数据库中没有）"""
    from pathlib import Path

    db = SessionLocal()
    try:
        # 检查是否已有任务书
        existing = db.query(TaskBook).count()
        if existing > 0:
            print(f"✓ 任务书已存在 {existing} 条，跳过初始化")
            return

        # GENERIC: 遍历 reference/ 目录下所有 .docx，不再硬编码文件名
        project_root = Path(__file__).parent.parent.parent
        task_books = []
        reference_dir = project_root / "reference"
        docx_paths = sorted(reference_dir.glob("*.docx")) if reference_dir.exists() else []
        if not docx_paths:
            print("⚠ 未找到外部任务书，使用仓库内置的脱敏演示模板")
            task_books = [
                TaskBook(
                    name="演示任务书（小学组）",
                    content_md=(
                        "# 演示任务书\n\n"
                        "本模板仅用于离线演示和界面验证，不代表任何真实赛项规则。\n\n"
                        "- 主题表达清楚\n"
                        "- 作品功能可说明\n"
                        "- 创作过程有记录\n"
                        "- 最终结论须由人工评审确认"
                    ),
                    group_type="小学组",
                    is_active=False,
                ),
                TaskBook(
                    name="演示任务书（初中组）",
                    content_md=(
                        "# 演示任务书\n\n"
                        "本模板仅用于离线演示和界面验证，不代表任何真实赛项规则。\n\n"
                        "- 主题表达清楚\n"
                        "- 作品功能可说明\n"
                        "- 创作过程有记录\n"
                        "- 最终结论须由人工评审确认"
                    ),
                    group_type="初中组",
                    is_active=False,
                ),
            ]

        for docx_path in docx_paths:
            # 从文件名推断 group_type
            stem_lower = docx_path.stem.lower()
            if "小学" in stem_lower or "primary" in stem_lower:
                group_type = "小学组"
            elif "初中" in stem_lower or "middle" in stem_lower:
                group_type = "初中组"
            else:
                group_type = "通用"

            # 解析Word为Markdown
            try:
                from docx import Document

                doc = Document(str(docx_path))
                lines = []

                # 解析段落
                for para in doc.paragraphs:
                    text = para.text.strip()
                    if not text:
                        continue
                    style = para.style.name.lower()
                    if 'heading 1' in style:
                        lines.append(f"# {text}")
                    elif 'heading 2' in style:
                        lines.append(f"## {text}")
                    elif 'heading 3' in style:
                        lines.append(f"### {text}")
                    elif 'heading 4' in style:
                        lines.append(f"#### {text}")
                    else:
                        lines.append(text)

                # 解析表格
                for table in doc.tables:
                    rows_md = []
                    for row in table.rows:
                        cells = [c.text.strip() for c in row.cells]
                        rows_md.append(" | ".join(cells))
                    if rows_md:
                        lines.append("\n")
                        lines.append("\n".join(rows_md))
                        lines.append("")

                content_md = "\n".join(lines)

                # 创建任务书（名称使用文件名，去掉扩展名）
                name = docx_path.stem
                tb = TaskBook(
                    name=name,
                    content_md=content_md,
                    group_type=group_type,
                    is_active=False,
                )
                task_books.append(tb)
                print(f"✓ 已解析任务书: {name}, Markdown长度: {len(content_md)}")

            except Exception as e:
                print(f"✗ 解析任务书失败 {docx_path}: {e}")
                continue

        if task_books:
            # 添加到数据库
            for tb in task_books:
                db.add(tb)

            db.commit()

            # 激活所有任务书
            db.query(TaskBook).update({TaskBook.is_active: False})
            for tb in task_books:
                tb.is_active = True

            db.commit()
            print(f"✓ 已初始化并激活 {len(task_books)} 个任务书")
        else:
            print("⚠ 没有成功解析的任务书")

    except Exception as e:
        print(f"✗ 初始化任务书失败: {e}")
        import traceback
        traceback.print_exc()
        db.rollback()
    finally:
        db.close()


def _init_default_competition():
    """GENERIC: 如果没有 competitions 记录，自动创建默认赛事"""
    db = SessionLocal()
    try:
        count = db.query(Competition).count()
        if count > 0:
            return

        default_dimensions = {
            "theme": {"name": "主题立意", "max_score": 20, "description": "思想性、原创性"},
            "presentation": {"name": "产品表现力", "max_score": 30, "description": "技术性、创新性、艺术性与用户体验"},
            "process": {"name": "过程完整性", "max_score": 30, "description": "过程文档、迭代思维、AIGC交互日志"},
            "ai_literacy": {"name": "人工智能素养", "max_score": 20, "description": "工具策略、批判性思维、伦理意识、沟通表达"},
        }

        default_comp = Competition(
            name="创意编程评分系统",
            description="通用创意编程作品评分系统",
            is_active=True,
            dimensions_config=default_dimensions,
            naming_pattern=r"^(.*)\+(.+)\.zip$",
            naming_example="小学组+张三.zip",
            judge_groups=4,
            judges_per_group=2,
            machine_weight=0.6,
            task_book_ids=[],
            scoring_mode="mixed",
            ai_weight=0.4,
            human_weight=0.6,
        )
        db.add(default_comp)
        db.commit()
        print("✓ 已创建默认赛事: 创意编程评分系统")
    except Exception as e:
        print(f"✗ 创建默认赛事失败: {e}")
        db.rollback()
    finally:
        db.close()


# ============ 队伍模型 ============
class Team(Base):
    """参赛队伍"""
    __tablename__ = "teams"

    id = Column(Integer, primary_key=True, autoincrement=True)
    team_code = Column(String(50), unique=True, nullable=False, index=True, comment="队伍编号（原始长编号）")
    short_code = Column(String(10), unique=True, nullable=False, index=True, comment="短编号（P001~P275）")
    team_name = Column(String(100), nullable=False, comment="团队名称")
    group_type = Column(String(20), nullable=False, comment="组别：小学组/初中组")
    school = Column(String(200), comment="学校名称")
    district = Column(String(100), comment="区县")
    teacher = Column(String(100), comment="指导老师（兼容旧数据）")
    teacher_name = Column(String(50), comment="指导老师姓名")
    teacher_phone = Column(String(20), comment="指导老师手机号")
    members = Column(JSON, comment="选手信息列表")
    judge_group = Column(Integer, comment="评审组号(1-4)")
    status = Column(String(20), default="confirmed", comment="状态：pending(待确认)/confirmed(已确认)")
    source = Column(String(20), default="excel", comment="来源：excel(导入)/work(作品生成)")
    created_at = Column(DateTime, default=datetime.now)
    updated_at = Column(DateTime, default=datetime.now, onupdate=datetime.now)

    # 关联
    work = relationship("Work", back_populates="team", uselist=False)
    machine_scores = relationship("MachineScore", back_populates="team")
    human_scores = relationship("HumanScore", back_populates="team")
    final_scores = relationship("FinalScore", back_populates="team")

    __table_args__ = (
        Index("idx_team_group", "group_type", "judge_group"),
    )


# ============ 作品模型 ============
class Work(Base):
    """参赛作品"""
    __tablename__ = "works"

    id = Column(Integer, primary_key=True, autoincrement=True)
    team_id = Column(Integer, ForeignKey("teams.id"), nullable=False, unique=True)
    original_filename = Column(String(500), comment="原始文件名")
    file_path = Column(String(500), comment="存储路径")
    file_size = Column(Integer, comment="文件大小(字节)")

    # 解析结果
    has_source = Column(Boolean, default=False, comment="是否包含源文件")
    has_aigc_log = Column(Boolean, default=False, comment="是否包含AIGC日志")
    has_screenshots = Column(Boolean, default=False, comment="是否包含截图/视频")
    has_readme = Column(Boolean, default=False, comment="是否包含说明文档")

    # 文件详情
    source_files = Column(JSON, comment="源文件列表")
    aigc_log_files = Column(JSON, comment="AIGC日志文件列表")
    screenshot_files = Column(JSON, comment="截图/视频文件列表")
    document_files = Column(JSON, comment="说明文档列表")  # 新增：文档文件
    readme_content = Column(Text, comment="README文档内容")  # 新增：文档文本

    # 代码分析结果
    code_language = Column(String(50), comment="编程语言")
    code_line_count = Column(Integer, comment="代码行数")
    comment_rate = Column(Float, comment="注释率")
    complexity_score = Column(Float, comment="复杂度评分")

    # AIGC日志分析
    aigc_tool_count = Column(Integer, comment="使用的AIGC工具数量")
    aigc_interaction_count = Column(Integer, comment="AIGC交互次数")
    aigc_tools = Column(JSON, comment="使用的AIGC工具列表")

    # 状态
    is_parsed = Column(Boolean, default=False, comment="是否已解析")
    is_flagged = Column(Boolean, default=False, comment="是否标记异常")
    flag_reason = Column(String(500), comment="异常标记原因")
    plagiarism_flag = Column(Boolean, default=False, comment="抄袭标记")

    # P0-3 (2026-05-22): 锚点等级 (P0-4 锚点分层使用)
    anchor_level = Column(String(10), default="Lv0", comment="锚点等级: Lv0/Lv1/Lv2/Lv3")

    created_at = Column(DateTime, default=datetime.now)
    updated_at = Column(DateTime, default=datetime.now, onupdate=datetime.now)

    team = relationship("Team", back_populates="work")


# ============ 机器评分模型 ============
class MachineScore(Base):
    """机器自动评分（多版本）"""
    __tablename__ = "machine_scores"

    id = Column(Integer, primary_key=True, autoincrement=True)
    team_id = Column(Integer, ForeignKey("teams.id"), nullable=False, index=True)
    team = relationship("Team", back_populates="machine_scores")

    # 评分版本信息
    model_name = Column(String(100), comment="使用的模型名称")
    is_adopted = Column(Boolean, default=False, comment="是否为最终采用的评分")
    scoring_session = Column(String(100), comment="评分批次ID（同一批次评分共享session）")

    # 各维度得分
    theme_score = Column(Float, default=0, comment="主题立意得分(0-20)")
    presentation_score = Column(Float, default=0, comment="产品表现力得分(0-30)")
    process_score = Column(Float, default=0, comment="过程完整性得分(0-30)")
    ai_literacy_score = Column(Float, default=0, comment="人工智能素养得分(0-20)")
    total_score = Column(Float, default=0, comment="机器评分总分(0-100)")

    # 评分详情
    scoring_details = Column(JSON, comment="评分详细分析")
    code_quality_details = Column(JSON, comment="代码质量详情")
    aigc_analysis_details = Column(JSON, comment="AIGC日志分析详情")
    feature_detection_details = Column(JSON, comment="功能检测详情")

    # AI评语
    theme_comment = Column(Text, comment="主题立意评语")
    presentation_comment = Column(Text, comment="产品表现力评语")
    process_comment = Column(Text, comment="过程完整性评语")
    ai_literacy_comment = Column(Text, comment="AI素养评语")
    overall_comment = Column(Text, comment="总评语")

    # 答辩问题
    defense_questions = Column(JSON, comment="答辩问题列表")

    # P0-3 (2026-05-22): 锚点置信度
    confidence = Column(String(20), default="UNKNOWN", comment="锚点置信度: HIGH/MEDIUM/LOW/UNKNOWN")

    # P0-5 (2026-05-22): Flag 规则引擎结果
    # 存 List[dict] (每条 dict 含 code/level/msg), SQLAlchemy JSON 自动序列化为 TEXT
    flags = Column(JSON, default=list, comment="Flag 列表: [{code,level,msg}, ...]")

    # P1-6 (2026-05-26): 校准后总分
    calibrated_score = Column(Float, default=0, comment="校准后总分(0-100)")

    # === 双轨评分：AI 维度分（Phase 1.1 新增）===
    # 混合模式下 AI 只评这 4 个维度，合计 40 分
    ai_theme_score = Column(Float, default=0, comment="主题基础契合 0-5")
    ai_code_quality_score = Column(Float, default=0, comment="代码质量 0-15")
    ai_completeness_score = Column(Float, default=0, comment="材料完整性 0-10")
    ai_aigc_score = Column(Float, default=0, comment="AIGC 规范性 0-10")
    ai_total_score = Column(Float, default=0, comment="AI 维度合计 0-40")

    # === 纯AI模式：AI 推断人工维度（Phase 2 新增）===
    inferred_presentation_score = Column(Float, default=0, comment="推断产品表现力 0-25")
    inferred_creativity_score = Column(Float, default=0, comment="推断创意深度 0-15")
    inferred_process_score = Column(Float, default=0, comment="推断过程深度 0-15")
    inferred_performance_score = Column(Float, default=0, comment="推断现场表现 0-5")
    inferred_total_score = Column(Float, default=0, comment="推断人工维度合计 0-60")

    # 状态
    is_completed = Column(Boolean, default=False, comment="是否评分完成")
    error_message = Column(Text, comment="错误信息")

    created_at = Column(DateTime, default=datetime.now)

    __table_args__ = (
        Index("idx_machine_team_model", "team_id", "model_name"),
        Index("idx_machine_team_adopted", "team_id", "is_adopted"),
    )


# ============ 人工评分模型 ============
class HumanScore(Base):
    """人工评分"""
    __tablename__ = "human_scores"

    id = Column(Integer, primary_key=True, autoincrement=True)
    team_id = Column(Integer, ForeignKey("teams.id"), nullable=False, index=True)
    judge_name = Column(String(100), nullable=False, comment="评委姓名")
    judge_group = Column(Integer, comment="评委所在组号")

    # 各维度得分
    theme_score = Column(Float, default=0, comment="主题立意得分(0-20)")
    presentation_score = Column(Float, default=0, comment="产品表现力得分(0-30)")
    process_score = Column(Float, default=0, comment="过程完整性得分(0-30)")
    ai_literacy_score = Column(Float, default=0, comment="人工智能素养得分(0-20)")
    total_score = Column(Float, default=0, comment="人工评分总分(0-100)")

    # 评语
    comment = Column(Text, comment="评委评语")

    # === 双轨评分：人工维度分（Phase 1.1 新增）===
    # 混合模式下评委现场输入这 4 个维度，合计 60 分
    human_presentation_score = Column(Float, default=0, comment="产品表现力 0-25")
    human_creativity_score = Column(Float, default=0, comment="创意深度 0-15")
    human_process_score = Column(Float, default=0, comment="过程深度 0-15")
    human_performance_score = Column(Float, default=0, comment="现场表现 0-5")
    human_total_score = Column(Float, default=0, comment="人工维度合计 0-60")

    created_at = Column(DateTime, default=datetime.now)
    updated_at = Column(DateTime, default=datetime.now, onupdate=datetime.now)

    team = relationship("Team", back_populates="human_scores")

    __table_args__ = (
        Index("idx_human_team_judge", "team_id", "judge_name"),
    )


# ============ 综合评分模型 ============
class FinalScore(Base):
    """综合评分（加权合成）"""
    __tablename__ = "final_scores"

    id = Column(Integer, primary_key=True, autoincrement=True)
    team_id = Column(Integer, ForeignKey("teams.id"), nullable=False, unique=True, index=True)

    # 机器评分
    machine_score = Column(Float, default=0, comment="机器评分总分")
    machine_weight = Column(Float, default=0.6, comment="机器评分权重")

    # 人工评分（多评委平均）
    human_score_avg = Column(Float, default=0, comment="人工评分平均分")
    human_judge_count = Column(Integer, default=0, comment="参与评分评委数")
    human_weight = Column(Float, default=0.4, comment="人工评分权重")

    # 综合得分
    final_score = Column(Float, default=0, comment="综合得分")
    ranking = Column(Integer, comment="排名")

    # === 双轨评分：综合分与模式记录（Phase 1.1 新增）===
    composite_score = Column(Float, default=0, comment="双轨综合分（混合：AI*0.4+人工*0.6）")
    scoring_mode = Column(String(20), default="mixed", comment="评分模式: mixed/ai_only")

    created_at = Column(DateTime, default=datetime.now)
    updated_at = Column(DateTime, default=datetime.now, onupdate=datetime.now)

    team = relationship("Team", back_populates="final_scores")


# ============ 比赛配置模型 ============
class Competition(Base):
    """比赛配置（每届比赛独立设置）"""
    __tablename__ = "competitions"

    id = Column(Integer, primary_key=True, autoincrement=True)
    name = Column(String(200), nullable=False, comment="比赛名称")
    year = Column(Integer, comment="年份")
    group_type = Column(String(20), comment="适用组别：小学组/初中组/通用")
    # 双轨评分模式：mixed（混合评分）/ ai_only（纯AI评分）
    scoring_mode = Column(String(20), default="mixed", comment="评分模式: mixed/ai_only")
    # AI vs 人工权重（支持微调，默认 AI 40% + 人工 60%）
    ai_weight = Column(Float, default=0.4, comment="AI 评分权重")
    human_weight = Column(Float, default=0.6, comment="人工评分权重")
    is_active = Column(Boolean, default=False, comment="是否为当前激活的比赛")
    description = Column(Text, comment="比赛描述")

    # GENERIC: Phase 2 新增字段
    dimensions_config = Column(JSON, default=dict, comment="评分维度配置")
    naming_pattern = Column(String(500), comment="作品命名正则")
    naming_example = Column(String(200), comment="命名示例")
    judge_groups = Column(Integer, default=4, comment="评审组数量")
    judges_per_group = Column(Integer, default=2, comment="每组评委数")
    machine_weight = Column(Float, default=0.6, comment="机器评分权重")
    task_book_ids = Column(JSON, default=list, comment="任务书ID列表")

    created_at = Column(DateTime, default=datetime.now)
    updated_at = Column(DateTime, default=datetime.now, onupdate=datetime.now)


# ============ 评委模型 ============
class Judge(Base):
    """评委"""
    __tablename__ = "judges"

    id = Column(Integer, primary_key=True, autoincrement=True)
    name = Column(String(100), nullable=False, unique=True, comment="评委姓名")
    password_hash = Column(String(200), nullable=False, comment="密码哈希")
    judge_group = Column(Integer, comment="所在评审组号")
    is_active = Column(Boolean, default=True, comment="是否活跃")

    created_at = Column(DateTime, default=datetime.now)


# ============ 操作日志模型 ============
class AuditLog(Base):
    """操作审计日志"""
    __tablename__ = "audit_logs"

    id = Column(Integer, primary_key=True, autoincrement=True)
    operator = Column(String(100), comment="操作人")
    action = Column(String(100), nullable=False, comment="操作类型")
    target_type = Column(String(50), comment="操作对象类型")
    target_id = Column(Integer, comment="操作对象ID")
    detail = Column(JSON, comment="操作详情")
    created_at = Column(DateTime, default=datetime.now)

    __table_args__ = (
        Index("idx_audit_time", "created_at"),
    )


# ============ 评审组配置模型 ============
class JudgeGroupConfig(Base):
    """评审组配置"""
    __tablename__ = "judge_group_configs"

    id = Column(Integer, primary_key=True, autoincrement=True)
    group_index = Column(Integer, nullable=False, comment="组序号(1,2,3...)")
    group_name = Column(String(100), nullable=False, comment="组名称")
    description = Column(String(500), comment="组描述")
    judges = Column(JSON, comment="评委列表")


# ============ 任务书模型 ============
class TaskBook(Base):
    """任务书"""
    __tablename__ = "task_books"

    id = Column(Integer, primary_key=True, autoincrement=True)
    name = Column(String(200), nullable=False, comment="任务书名称")
    group_type = Column(String(20), comment="适用组别：小学组/初中组/通用")
    content_md = Column(Text, comment="任务书内容 Markdown 格式")
    is_active = Column(Boolean, default=False, comment="是否激活使用")
    created_at = Column(DateTime, default=datetime.now)
    updated_at = Column(DateTime, default=datetime.now, onupdate=datetime.now)


# ============ 评审表模型 ============
class ScoringForm(Base):
    """评审表"""
    __tablename__ = "scoring_forms"

    id = Column(Integer, primary_key=True, autoincrement=True)
    name = Column(String(200), nullable=False, comment="评审表名称")
    dimensions = Column(JSON, comment="维度配置JSON数组")
    is_default = Column(Boolean, default=False, comment="是否系统默认")
    is_active = Column(Boolean, default=True, comment="是否激活")
    created_at = Column(DateTime, default=datetime.now)
    updated_at = Column(DateTime, default=datetime.now, onupdate=datetime.now)
