"""Task-scoped exports. Never fall back to legacy weighted database scores."""
from __future__ import annotations

import csv
import io
import json
import sys
from dataclasses import dataclass
from typing import Any, Literal
from uuid import UUID

from services.result_derivation_service import DerivedResult, ResultDerivationError

ExportFormat = Literal["xlsx", "docx", "csv"]


class AuthoritativeExportError(Exception):
    def __init__(self, code: str, message: str, status: int = 409, blocked_items=None):
        super().__init__(message)
        self.code, self.message, self.status = code, message, status
        self.blocked_items = blocked_items or []


def require_task_id(task_id: str | None) -> str:
    if not task_id:
        raise AuthoritativeExportError(
            "EXPORT_TASK_REQUIRED", "请在结果导出页选择评分任务；旧综合分不能直接作为权威结果导出。"
        )
    try:
        if str(UUID(task_id)) != task_id:
            raise ValueError
    except (ValueError, AttributeError):
        raise AuthoritativeExportError("EXPORT_TASK_INVALID", "评分任务编号无效。", 400)
    return task_id


def result_payload(result: DerivedResult) -> dict:
    return {
        "item_id": result.item_id,
        "exportable": True,
        "authority_type": result.authority_type,
        "derivation_id": result.derivation_id,
        "adoption_id": result.adoption_id,
        "lock_id": result.lock_id,
        "result": result.final_snapshot,
    }


@dataclass(frozen=True)
class ExportSelection:
    task_id: str
    task_revision: int
    total_items: int
    results: tuple[DerivedResult, ...]


class AuthoritativeExportService:
    def __init__(self, manager: Any, derivation: Any):
        self.manager = manager
        self.derivation = derivation

    def list_tasks(self) -> list[dict]:
        rows, offset = [], 0
        while True:
            tasks = self.manager.list_tasks(offset=offset, limit=200)
            rows.extend({
                "task_id": task.task_id, "batch_id": task.batch_id,
                "total_items": task.total_items, "status": task.status,
            } for task in tasks if task.task_type == "scoring_pipeline")
            if len(tasks) < 200:
                return rows
            offset += 200

    def _task_items(self, task_id: str):
        require_task_id(task_id)
        task = self.manager.get_task(task_id)
        if task is None or task.task_type != "scoring_pipeline":
            raise AuthoritativeExportError("EXPORT_TASK_NOT_FOUND", "未找到评分任务。", 404)
        items = self.manager.list_items(task_id)
        expected = [entry.item_id for entry in task.item_index]
        actual = [item.item_id for item in items]
        if (len(expected) != task.total_items or len(set(expected)) != len(expected)
                or len(actual) != len(expected) or set(actual) != set(expected)
                or any(item.task_id != task_id or item.task_type != "scoring_pipeline" for item in items)):
            raise AuthoritativeExportError("EXPORT_TASK_INCOMPLETE", "任务条目记录不完整，已阻止导出。")
        return task, sorted(items, key=lambda item: item.item_id)

    def preview(self, task_id: str) -> dict:
        task, items = self._task_items(task_id)
        rows = []
        for item in items:
            try:
                rows.append(result_payload(self.derivation.derive(task_id, item.item_id)))
            except ResultDerivationError as exc:
                rows.append({
                    "item_id": item.item_id, "exportable": False,
                    "error": {"code": exc.error_code, "reasons": list(exc.reasons)},
                })
        return {"task_id": task_id, "total_items": task.total_items, "items": rows}

    def collect(self, task_id: str, item_ids: list[str] | None = None) -> ExportSelection:
        task, items = self._task_items(task_id)
        known = {item.item_id for item in items}
        selected = sorted(known) if item_ids is None else item_ids
        if not selected or len(set(selected)) != len(selected) or not set(selected).issubset(known):
            raise AuthoritativeExportError("EXPORT_SELECTION_INVALID", "请选择该任务中不重复的有效条目。", 400)
        results, blocked = [], []
        for item_id in sorted(selected):
            try:
                result = self.derivation.derive(task_id, item_id)
                if result.task_id != task_id or result.item_id != item_id:
                    raise AuthoritativeExportError("EXPORT_RESULT_MISMATCH", "结果与所选条目不一致。")
                results.append(result)
            except ResultDerivationError as exc:
                blocked.append({"item_id": item_id, "code": exc.error_code, "reasons": list(exc.reasons)})
        if blocked:
            raise AuthoritativeExportError(
                "EXPORT_BLOCKED_BY_REVIEW", "所选条目存在未确认或被阻断的结果，请处理后重试。", blocked_items=blocked
            )
        return ExportSelection(task_id, task.revision, task.total_items, tuple(results))

    def export(self, task_id: str, fmt: ExportFormat, item_ids: list[str] | None = None,
               expected_derivations: dict[str, str] | None = None) -> bytes:
        if fmt not in ("xlsx", "docx", "csv"):
            raise AuthoritativeExportError("EXPORT_FORMAT_INVALID", "不支持的导出格式。", 400)
        selection = self.collect(task_id, item_ids)
        if expected_derivations is not None and expected_derivations != {r.item_id: r.derivation_id for r in selection.results}:
            raise AuthoritativeExportError("EXPORT_RESULT_CHANGED", "页面中的结果已变化，请刷新后重新选择。")
        payload = {"xlsx": _xlsx, "docx": _docx, "csv": _csv}[fmt](selection)
        # Rendering may take time. Recheck every selected authority before releasing any bytes.
        current = self.collect(task_id, [result.item_id for result in selection.results])
        if (current.task_revision != selection.task_revision or current.total_items != selection.total_items
                or [r.derivation_id for r in current.results] != [r.derivation_id for r in selection.results]):
            raise AuthoritativeExportError("EXPORT_RESULT_CHANGED", "生成期间结果发生变化，请刷新后重新导出。")
        return payload


HEADERS = ["条目编号", "确认方式", "总分", "客观分", "主观分", "材料编号", "评分任务", "结果版本", "评分尝试", "推导编号", "采用记录", "锁定记录"]


def _row(result: DerivedResult) -> list:
    return [result.item_id, result.authority_type, result.total_score, result.objective_score,
            result.subjective_score, result.submission_id, result.task_id, result.snapshot_id,
            result.attempt_id, result.derivation_id, result.adoption_id, result.lock_id]


def _csv(selection: ExportSelection) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(HEADERS + ["维度分", "导出条数", "任务总条数"])
    for result in selection.results:
        values = _row(result) + [json.dumps(result.final_snapshot["dimension_scores"], ensure_ascii=False), len(selection.results), selection.total_items]
        writer.writerow(["'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r", "\n")) else value for value in values])
    return stream.getvalue().encode("utf-8-sig")


def _xlsx(selection: ExportSelection) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    notes = book.active
    notes.title = "导出说明"
    notes.append(["项目", "内容"])
    for row in [
        ["评分任务", selection.task_id], ["导出条数", len(selection.results)], ["任务总条数", selection.total_items],
        ["范围", "仅包含本次明确选择的条目，不代表全部任务或历史队伍。"],
        ["分数来源", "已采用或锁定的评分任务结果，不使用旧版综合分加权。"],
        ["缺失值", "空白表示该快照未提供该类分数，不代表零分。"],
        ["确认方式", "manual_final_lock：人工最终锁定；result_adoption：已采用结果。"],
    ]:
        notes.append(row)
    scores = book.create_sheet("权威结果")
    scores.append(HEADERS)
    dimensions = book.create_sheet("维度分")
    dimensions.append(["条目编号", "维度", "得分", "最低分", "最高分", "分值范围引用"])
    for result in selection.results:
        scores.append(_row(result))
        for dimension in result.dimension_scores:
            dimensions.append([result.item_id, dimension["dimension_code"], dimension["score"], dimension["min_score"], dimension["max_score"], dimension["score_range_ref"]])
    widths = [[22, 86], [40, 24, 12, 12, 12, 40, 40, 38, 38, 45, 38, 38], [40, 24, 12, 12, 12, 30]]
    for sheet, column_widths in zip(book.worksheets, widths):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.font = Font(name="Microsoft YaHei", size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "General"
        for cell in sheet[1]:
            cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24435A")
        for index, width in enumerate(column_widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for row_number in range(2, sheet.max_row + 1):
            sheet.row_dimensions[row_number].height = 48
        sheet.print_title_rows = "1:1"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def _docx(selection: ExportSelection) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    document = Document()
    cjk_font = "Heiti SC" if sys.platform == "darwin" else "Microsoft YaHei" if sys.platform == "win32" else "Noto Sans CJK SC"
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(18)
    section.left_margin = section.right_margin = Mm(20)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = cjk_font
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().rFonts
        for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{attribute}"), None)
        fonts.set(qn("w:eastAsia"), cjk_font)
    document.styles["Normal"].font.size = Pt(10)
    document.styles["Normal"].paragraph_format.space_after = Pt(4)
    for index, result in enumerate(selection.results):
        if index:
            document.add_page_break()
        document.add_heading("评审权威结果", 0)
        document.add_paragraph(f"本次导出{len(selection.results)}条，任务共{selection.total_items}条。当前为第{index + 1}条。")
        document.add_paragraph("分数来自已采用或锁定的评分结果；未使用旧版综合分加权。空白分数表示未提供，不代表零分。")
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        table.autofit = False
        table.columns[0].width, table.columns[1].width = Mm(28), Mm(142)
        for label, value in zip(HEADERS, _row(result)):
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = "" if value is None else str(value)
            cells[0].width, cells[1].width = Mm(28), Mm(142)
        document.add_heading("维度评分", level=2)
        dimensions = document.add_table(rows=1, cols=4)
        dimensions.style = "Table Grid"
        for cell, label in zip(dimensions.rows[0].cells, ["维度", "得分", "最低分", "最高分"]):
            cell.text = label
        for dimension in result.dimension_scores:
            for cell, key in zip(dimensions.add_row().cells, ["dimension_code", "score", "min_score", "max_score"]):
                cell.text = str(dimension[key])
        for tbl in (table, dimensions):
            for row in tbl.rows:
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()
